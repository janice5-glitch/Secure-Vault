import sys, json
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PAGES = json.load(open("toc_pages.json")) if len(sys.argv) > 1 else {}
NAME, ROLL, REG = "Janice A.R", "25IT1145", "312425205079"
TITLE = "SecureVault: A Secure Password Manager using AES-256 Encryption"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
for m in ("left_margin", "right_margin"): setattr(sec, m, Inches(1))
sec.top_margin = sec.bottom_margin = Inches(0.9)
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(4)

def para(text="", bold=False, italic=False, size=None, align=None, after=None, before=None, font=None, color=None):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        if size: r.font.size = Pt(size)
        if font: r.font.name = font
        if color: r.font.color.rgb = RGBColor(*color)
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if before is not None: p.paragraph_format.space_before = Pt(before)
    return p

def heading(n, text, newpage=False):
    p = para(f"{n}. {text}:", bold=True, size=14, before=6, after=8)
    if newpage: p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True
    return p

def sub(text):
    p = para(text, bold=True, before=6, after=2); p.paragraph_format.keep_with_next = True

def body(text):
    p = para(text, after=6); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.line_spacing = 1.3; return p

def bullet(text, boldlead=None):
    p = doc.add_paragraph(style="List Bullet")
    if boldlead:
        r = p.add_run(boldlead); r.bold = True
    p.add_run(text); p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.2
    return p

def page_field(par):
    r = par.add_run()
    for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if t:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
        r._r.append(e)

def borders(table):
    tbl = table._tbl; pr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:color"), "000000"); b.append(e)
    pr.append(b)

def code(text):
    for line in text.strip("\n").split("\n"):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Inches(0.15)
        r = p.add_run(line if line else " "); r.font.name = "Courier New"; r.font.size = Pt(8.5)
        r._r.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")

def shot(file, caption, w=5.5):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(f"shots/{file}", width=Inches(w))
    c = para(caption, italic=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

# ------------------------------------------------------------------ COVER
hdr = doc.add_paragraph(); hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
hdr.add_run().add_picture("docs/college_header.png", width=Inches(6.2))
for _ in range(5): para()
para(TITLE, bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
para("MINI PROJECT REPORT", bold=True, italic=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, after=40)
for _ in range(2): para()
para("Submitted By", size=13, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
para(f"{NAME} ( {REG} )", bold=True, size=17, align=WD_ALIGN_PARAGRAPH.CENTER, font="Calibri", after=4)
para(f"Roll No: {ROLL}", bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER, font="Calibri", after=50)
for _ in range(3): para()
para("II Year", bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
para("B. Tech Information Technology", bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER)

# ------------------------------------------------------------------ TOC
doc.add_page_break()
para("TABLE OF CONTENTS", size=13, align=WD_ALIGN_PARAGRAPH.CENTER, after=16)
SECTIONS = ["Introduction", "Problem Statement", "Scope/Application of the Project", "Advantages of the Project",
            "System Requirements", "Project Modules", "Working of the Project", "Algorithms Used", "Project Code",
            "Output Screenshot", "Conclusion", "Future Enhancements", "References"]
t = doc.add_table(rows=1, cols=3); borders(t); t.autofit = False
widths = (Inches(0.8), Inches(4.4), Inches(1.1))
for i, h in enumerate(("Sl. No", "Contents", "Page No")):
    c = t.rows[0].cells[i]; c.width = widths[i]; c.text = ""
    r = c.paragraphs[0].add_run(h); r.bold = True; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
for i, s in enumerate(SECTIONS, 1):
    row = t.add_row().cells
    vals = (f"{i}.", s, str(PAGES.get(s, "")))
    for j, v in enumerate(vals):
        row[j].width = widths[j]; row[j].text = ""
        pp = row[j].paragraphs[0]; pp.add_run(v)
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER if j != 1 else WD_ALIGN_PARAGRAPH.LEFT
for r in t.rows:
    r.height = Inches(0.55)
    for c in r.cells:
        tcPr = c._tc.get_or_add_tcPr(); va = OxmlElement("w:vAlign"); va.set(qn("w:val"), "center"); tcPr.append(va)

# new section: page numbering restarts at 1
ns = doc.add_section(WD_SECTION.NEW_PAGE)
ns.footer.is_linked_to_previous = False
fp = ns.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; page_field(fp)
pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1"); ns._sectPr.append(pg)

# ------------------------------------------------------------------ 1
heading(1, "Introduction")
body("In the digital age every person maintains dozens of online accounts - email, banking, social media, college portals and cloud services. "
     "Remembering a unique, complex password for each of them is practically impossible, so people reuse simple passwords or write them down in "
     "insecure places. Such habits are the root cause of most account-takeover attacks, credential-stuffing attacks and data breaches.")
body("SecureVault is a locally hosted Password Manager that solves this problem. The user only has to remember one strong Master Password. "
     "All other credentials (website, username, password and notes) are stored inside an encrypted vault file. The vault is protected using "
     "industry-standard cryptography: the Master Password is converted into a 256-bit key using PBKDF2-HMAC-SHA256 with 600,000 iterations and a "
     "random salt, and the entire vault is encrypted with AES-256 in Galois/Counter Mode (AES-GCM), which provides both confidentiality and "
     "tamper detection.")
body("The application is written in Python using the Flask web framework and the 'cryptography' library. It provides a clean browser based "
     "interface for adding, viewing, editing, searching and deleting credentials, a cryptographically secure password generator, and a live password "
     "strength meter based on entropy. Security features such as auto-lock after inactivity, login attempt lock-out and clipboard clearing "
     "further protect the user.")

heading(2, "Problem Statement")
body("Most users either reuse the same weak password across many websites or store passwords in plain text files, notebooks, browser notes or "
     "messaging apps. When one website is breached, all accounts that share that password become vulnerable. Many existing cloud-based password "
     "managers also require trusting a third party with sensitive data. There is a need for a simple, transparent and secure tool that "
     "(i) stores credentials in an encrypted form, (ii) unlocks only with one Master Password, (iii) generates strong random passwords, "
     "(iv) evaluates password strength, and (v) keeps all data under the user's own control without depending on any external server.")

# ------------------------------------------------------------------ 3
heading(3, "Scope/Application of the Project", newpage=True)
for title, pts in [
    ("Individual Users", ["Safely stores passwords of email, banking, social media and shopping accounts.",
                          "Removes the need to remember or reuse passwords.",
                          "Generates strong unique passwords for every new account."]),
    ("Educational Institutions", ["Students and faculty can keep portal, LMS and research-account credentials securely.",
                                  "Can be used as a teaching example for applied cryptography and secure software design."]),
    ("Small Businesses and Teams", ["Provides a lightweight, self-hosted vault for storing shared service credentials.",
                                    "No subscription cost and no data leaves the organisation."]),
    ("Developers and IT Administrators", ["Stores API keys, server logins and database credentials in encrypted form.",
                                          "Source code is open and extensible with additional modules."]),
    ("Security Awareness", ["The dashboard highlights weak passwords, encouraging users to improve their security hygiene.",
                            "Demonstrates concepts such as key derivation, authenticated encryption and brute-force protection."]),
]:
    p = para(title, after=1, before=6); p.runs[0].bold = False
    for x in pts: bullet(x)

heading(4, "Advantages of the Project", newpage=True)
for t_, d in [("Strong Encryption", "The vault is encrypted with AES-256-GCM. Without the Master Password the file is unreadable, and any modification of the file is detected and rejected."),
              ("Single Master Password", "The user remembers only one password. The key is derived with PBKDF2 (600,000 iterations), which makes brute-force and dictionary attacks very slow."),
              ("Offline and Private", "All data is stored locally in one file. No cloud server, third party or internet connection is needed, which removes the risk of a central data breach."),
              ("Secure Password Generator", "Random passwords are generated with Python's 'secrets' module, a cryptographically secure random number generator, with user-selectable length and character types."),
              ("Password Strength Analysis", "A live entropy-based meter rates each password from Very Weak to Very Strong, suggests improvements and flags common passwords."),
              ("Auto-Lock and Lock-out", "The vault locks automatically after 5 minutes of inactivity, and five wrong master-password attempts trigger a temporary lock-out."),
              ("Easy to Use and Low Cost", "A clean web interface with search, show/hide and copy buttons. Built entirely using free, open-source software.")]:
    sub(f"{t_}")
    p = body(d)

# ------------------------------------------------------------------ 5
heading(5, "System Requirements", newpage=True)
sub("Hardware Requirements")
for x in ("Processor: Intel Core i3 or higher (or equivalent)", "RAM: 4 GB minimum", "Storage: 100 MB free disk space", "Display, keyboard and mouse"): bullet(x)
sub("Software Requirements")
for b, x in [("Operating System: ", "Windows 10/11, Linux or macOS"), ("Language: ", "Python 3.9 or higher"),
             ("Web Framework: ", "Flask 3.x"), ("Cryptography Library: ", "cryptography 42+ (AES-GCM, PBKDF2)"),
             ("Front End: ", "HTML5, CSS3 and JavaScript (Fetch API)"), ("Browser: ", "Google Chrome, Microsoft Edge or Firefox"),
             ("Editor / IDE: ", "Visual Studio Code or PyCharm"), ("Testing: ", "Python unittest, Playwright (for screenshots)")]: bullet(x, b)
sub("Installation Commands")
code("pip install -r requirements.txt\npython app.py\n# open http://127.0.0.1:5000 in the browser")

# ------------------------------------------------------------------ 6
heading(6, "Project Modules", newpage=True)
for n, (t_, pur, comp) in enumerate([
    ("Master Password and Authentication Module", "Creates the vault and verifies the user before decrypting it.",
     "Login/Setup page, PBKDF2 key derivation, attempt counter with 60-second lock-out, session token."),
    ("Encryption / Decryption Module (vault.py)", "Encrypts and decrypts the entire credential store.",
     "AES-256-GCM, random 12-byte nonce, random 16-byte salt, authentication tag for tamper detection."),
    ("Credential Management Module", "Add, view, edit, delete and search stored credentials.",
     "Dashboard table, Add/Edit form, unique entry IDs, timestamps, search by site or username."),
    ("Password Generator Module", "Creates strong random passwords on demand.",
     "secrets.SystemRandom, length slider (8-40), uppercase/digits/symbols options."),
    ("Password Strength Checker Module", "Evaluates password quality and gives improvement tips.",
     "Entropy calculation, common-password blacklist, five-level meter shown live."),
    ("Session and Security Module", "Protects the unlocked vault in use.",
     "Auto-lock after 5 minutes, manual lock/logout, clipboard clearing after 20 seconds, hidden-by-default passwords."),
    ("Testing Module (test_vault.py)", "Verifies correctness and security properties.",
     "7 unit tests: round-trip, wrong password, file encryption, tamper detection, CRUD, generator, strength."),
], 1):
    sub(f"{n}. {t_}")
    bullet(" " + pur, "Purpose:"); bullet(" " + comp, "Components:")

# ------------------------------------------------------------------ 7
heading(7, "Working of the Project", newpage=True)
body("When the application is started for the first time, the user is asked to create a Master Password. The password must be at least eight "
     "characters long and pass the strength check. A random 16-byte salt is generated, and the Master Password and salt are fed into PBKDF2-HMAC-SHA256 "
     "(600,000 iterations) to derive a 256-bit key. An empty credential list is encrypted using AES-256-GCM with a fresh random 12-byte nonce, and the "
     "salt, nonce and ciphertext are saved in the file vault.dat. The Master Password itself and the key are never written to disk.")
body("On later visits, the user enters the Master Password. The key is derived again using the stored salt and the vault is decrypted. If the password "
     "is wrong - or the file has been modified - the GCM authentication tag does not match, decryption fails and access is denied. After five "
     "failed attempts the login is blocked for 60 seconds. After a successful login, the decrypted list of entries is held only in memory and a "
     "random session token is issued to the browser.")
body("In the dashboard the user can add a new credential, search by website or username, show/hide and copy a password, edit or delete an entry. "
     "Every change re-encrypts the whole vault with a new nonce and saves it. While typing a password, the strength meter calls the server and "
     "displays the entropy and a rating; the Generate button creates a secure random password. If the user is inactive for five minutes or presses "
     "Lock, the decrypted data and key are erased from memory and the login page is shown again.")
sub("Process Flow")
code("""
 User --> Master Password --> PBKDF2-HMAC-SHA256 (salt, 600k) --> 256-bit key
                                                                     |
 vault.dat (salt | nonce | ciphertext) --> AES-256-GCM decrypt <----+
                                              |-- tag mismatch --> "Incorrect master password"
                                              v
                      In-memory entries --> Add / Edit / Delete / Search / Generate
                                              |
                      Every change --> new nonce --> AES-256-GCM encrypt --> save vault.dat
""")

# ------------------------------------------------------------------ 8
heading(8, "Algorithms Used", newpage=True)
body("The project relies on well-established cryptographic algorithms. No custom or home-made cryptography is used.")
for t_, typ, d in [
    ("Key Derivation - PBKDF2-HMAC-SHA256", "Password-Based Key Derivation Function 2",
     "Converts the human-readable Master Password into a 256-bit key by applying HMAC-SHA256 600,000 times with a random 16-byte salt. The large iteration count makes each guess expensive, slowing down brute-force attacks, and the unique salt defeats rainbow-table attacks."),
    ("Encryption - AES-256-GCM", "Authenticated Encryption with Associated Data (AEAD)",
     "Advanced Encryption Standard with a 256-bit key in Galois/Counter Mode. It encrypts the vault and also produces a 128-bit authentication tag. A new random 12-byte nonce is used for every save, so identical data never produces identical ciphertext."),
    ("Random Password Generation - CSPRNG", "Cryptographically Secure Pseudo-Random Number Generator",
     "Uses secrets.choice and secrets.SystemRandom().shuffle (backed by the operating system's entropy source). One character is guaranteed from every selected set, the rest are picked randomly and the final order is shuffled."),
    ("Password Strength - Shannon Entropy Estimation", "Entropy = length x log2(character pool size)",
     "The pool size is the sum of the character classes used (26 + 26 + 10 + 32). Entropy below 28 bits is Very Weak and above 80 bits is Very Strong; passwords found in a common-password list are always rated Very Weak."),
    ("Brute-Force Protection - Attempt Limiting", "Rate limiting / temporary lock-out",
     "A counter tracks wrong master-password attempts; after five failures the login is disabled for 60 seconds."),
]:
    sub(t_); bullet(" " + typ, "Algorithm Type:"); bullet(" " + d)

# ------------------------------------------------------------------ 9
heading(9, "Project Code", newpage=True)
sub("File: vault.py (cryptographic core)")
code(open("vault.py").read())
para()
sub("File: app.py (Flask web application)")
code(open("app.py").read())
para()
sub("File: test_vault.py (unit tests)")
code(open("test_vault.py").read())
body("The HTML templates (login, dashboard, add/edit and generator pages) and the CSS stylesheet are supplied in the project folder "
     "(templates/ and static/style.css).")

# ------------------------------------------------------------------ 10
heading(10, "Output Screenshot", newpage=True)
shots = [("01_setup.png", "Fig 10.1 - Vault creation with master password and live strength meter"),
         ("02_wrong_password.png", "Fig 10.2 - Login rejected when the master password is wrong"),
         ("03_add_entry.png", "Fig 10.3 - Adding a new credential with password generation"),
         ("04_dashboard.png", "Fig 10.4 - Dashboard with saved credentials, strength tags and statistics"),
         ("05_show_password.png", "Fig 10.5 - Revealing a password using the Show button"),
         ("06_search.png", "Fig 10.6 - Searching the vault"),
         ("07_generator.png", "Fig 10.7 - Secure password generator with options and entropy"),
         ("08_locked.png", "Fig 10.8 - Vault locked; unlock screen shown again"),
         ("09_unit_tests.png", "Fig 10.9 - All 7 unit tests passing"),
         ("10_encrypted_vault.png", "Fig 10.10 - Contents of vault.dat: only AES-256-GCM ciphertext is stored")]
for i, (f, c) in enumerate(shots):
    shot(f, c, 5.3 if i < 8 else 6.0)

# ------------------------------------------------------------------ 11,12,13
heading(11, "Conclusion", newpage=True)
body("SecureVault demonstrates how a password manager can be built securely and simply using proven cryptographic primitives. By combining "
     "PBKDF2 key derivation with AES-256-GCM authenticated encryption, the application guarantees that stored credentials remain confidential and "
     "that any tampering with the vault is detected. The secure generator and the entropy-based strength meter encourage users to adopt unique, "
     "strong passwords, while auto-lock, lock-out and clipboard clearing reduce the risk of misuse. The unit tests confirm that the encryption, "
     "decryption, tamper detection and password logic work correctly. The project therefore meets its objective of giving users a private, "
     "low-cost and easy to use solution to protect their digital identities.")
heading(12, "Future Enhancements")
body("In the future the project can be extended with several features. Firstly, two-factor authentication (TOTP) and biometric unlock can be added "
     "for the master login. Secondly, a browser extension with auto-fill would allow credentials to be entered directly on websites. Thirdly, "
     "encrypted cloud backup and multi-device synchronisation could be added using end-to-end encryption. Another enhancement is a breach "
     "checker that uses the Have I Been Pwned k-anonymity API to warn users about leaked passwords. Additional improvements include password expiry "
     "reminders, secure sharing of entries with team members, import/export of credentials in encrypted form, a mobile application and use of the "
     "stronger Argon2id key-derivation function.")
heading(13, "References")
for i, r in enumerate(["https://cryptography.io/en/latest/hazmat/primitives/aead/ - AES-GCM (cryptography library)",
                       "https://cryptography.io/en/latest/hazmat/primitives/key-derivation-functions/ - PBKDF2HMAC",
                       "https://flask.palletsprojects.com/ - Flask documentation",
                       "https://docs.python.org/3/library/secrets.html - Python secrets module",
                       "https://csrc.nist.gov/pubs/sp/800/38/d/final - NIST SP 800-38D: Galois/Counter Mode (GCM)",
                       "https://pages.nist.gov/800-63-3/sp800-63b.html - NIST SP 800-63B: Digital Identity Guidelines",
                       "https://cheatsheetseries.owasp.org/cheatsheets/Password_Storage_Cheat_Sheet.html - OWASP Password Storage Cheat Sheet"], 1):
    para(f"{i}. {r}", after=4)

doc.save("SecureVault_Mini_Project_Report.docx")
